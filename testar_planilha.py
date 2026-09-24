import os
import django

from openpyxl import load_workbook
from datetime import datetime, time

os.environ.setdefault("DJANGO_SETTINGS_MODULE", "Projeto_Integrador_IV.settings")
django.setup()

from ItaqueriAero.models import LocalDeVoo, HistoricoDiario


def converter_hora(valor):
    if valor is None:
        return None

    if isinstance(valor, time):
        return valor

    return datetime.strptime(valor, "%H:%M").time()


arquivo = "historico_rampa_completo_v5.xlsx"

planilha = load_workbook(arquivo)
aba = planilha["Dados diários"]

cabecalho = [celula.value for celula in aba[4]]

local = LocalDeVoo.objects.get(id=1)

criados = 0
atualizados = 0

for linha in aba.iter_rows(min_row=5, values_only=True):

    if all(valor is None for valor in linha):
        continue

    dados = dict(zip(cabecalho, linha))

    data = dados["data"].date()

    registro, criado = HistoricoDiario.objects.update_or_create(
        local=local,
        data=data,
        fonte=dados["fonte"],
        defaults={
            "temperatura_media_c": dados["temperatura_media_c"],
            "temperatura_maxima_c": dados["temperatura_maxima_c"],
            "hora_maxima": converter_hora(dados["hora_maxima"]),
            "temperatura_minima_c": dados["temperatura_minima_c"],
            "hora_minima": converter_hora(dados["hora_minima"]),
            "chuva_mm": dados["chuva_mm"],
            "umidade_media_pct": dados["umidade_media_pct"],
            "vento_medio_kmh": dados["vento_medio_kmh"],
            "vento_maximo_kmh": dados["vento_maximo_kmh"],
            "hora_vento_maximo": converter_hora(
                dados["hora_vento_maximo"]
            ),
            "direcao_vento_graus": dados["direcao_vento_graus"],
            "url_origem": dados["url_origem"],
            "observacao": dados["observacao"],
        },
    )

    if criado:
        criados += 1
    else:
        atualizados += 1


print("Importação concluída!")
print("Criados:", criados)
print("Atualizados:", atualizados)
print("Total no banco:", HistoricoDiario.objects.count())